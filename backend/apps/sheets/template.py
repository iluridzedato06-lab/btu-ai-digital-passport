"""Builds the official Google-Sheet-compatible template workbook and the local demo sheet."""
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import schema as S

HEADER_FILL = PatternFill("solid", fgColor="111111")
SYS_FILL = PatternFill("solid", fgColor="BDBDBD")
README_ROWS = [
    ("BTU AI Skills Digital Passport — data sheet", ""),
    ("ბიზნესისა და ტექნოლოგიების უნივერსიტეტი — AI უნარების ციფრული პასპორტი", ""),
    ("", ""),
    ("How it works (EN)", "The passport system reads this spreadsheet every minute. Each confirmed row in "
     "the Awards tab adds a stamp to the student's passport and issues a new, blockchain-anchored version."),
    ("როგორ მუშაობს (KA)", "სისტემა ამ ცხრილს ყოველ წუთს კითხულობს. Awards ჩანართში ყოველი დადასტურებული "
     "ჩანაწერი სტუდენტის პასპორტს ახალ შტამპს ამატებს და ქმნის ბლოკჩეინზე დადასტურებულ ახალ ვერსიას."),
    ("", ""),
    ("Students tab", "One row per student. Student fills names and BTU email; consent must be YES. "
     "Names (EN) are locked after the first passport is issued."),
    ("Students ჩანართი", "თითო სტრიქონი თითო სტუდენტზე. სახელი და გვარი ლათინურად, ელფოსტა @btu.edu.ge, "
     "consent = YES. პირველი პასპორტის გაცემის შემდეგ სახელის შეცვლა შეუძლებელია."),
    ("SkillCatalog tab", "Maintained by BTU staff. stamp_type: FOCUSED (AI Skill), APPLIED (AI Applied "
     "Skill) or ADVANCED (AI Advanced Skill)."),
    ("Awards tab", "One row per earned skill. award_id must be unique (e.g. AW-0001). Only rows with "
     "confirmed = TRUE and confirmed_by filled in are added to passports. date_earned uses YYYY-MM-DD."),
    ("Awards ჩანართი", "თითო სტრიქონი თითო მოპოვებულ უნარზე. პასპორტში ემატება მხოლოდ confirmed = TRUE "
     "ჩანაწერები, სადაც confirmed_by შევსებულია. თარიღის ფორმატი: YYYY-MM-DD."),
    ("Grey sys_* columns", "Written by the system (status, passport ID, version). Do not edit them."),
    ("ნაცრისფერი sys_* სვეტები", "ივსება სისტემის მიერ. არ შეცვალოთ."),
    ("", ""),
    ("Error codes", "; ".join(f"{k} {v.split('{')[0].strip()}" for k, v in S.ERRORS.items())),
]

CATALOG_ROWS = [
    ["PROMPT-ENG", "Prompt Engineering", "პრომპტ ინჟინერია", "FOCUSED",
     "Ability to design, structure and refine prompts to effectively interact with generative AI systems.",
     "გენერაციულ AI სისტემებთან ეფექტური მუშაობისთვის პრომპტების შექმნის, სტრუქტურირებისა და დახვეწის უნარი.",
     "Prompt Engineering Workshop", "YES"],
    ["AI-WORKFLOW", "AI Workflow Automation", "სამუშაო პროცესების ავტომატიზაცია AI-ით", "APPLIED",
     "Ability to design and implement AI-enabled workflows to automate tasks and improve efficiency.",
     "AI-ზე დაფუძნებული სამუშაო პროცესების დაგეგმვისა და დანერგვის უნარი ამოცანების ავტომატიზაციისა და ეფექტიანობისთვის.",
     "AI Automation Bootcamp (3-day intensive)", "YES"],
    ["DATA-AI", "Data Analysis with AI", "მონაცემთა ანალიზი AI-ით", "ADVANCED",
     "Ability to apply AI-enabled tools to collect, analyze, visualize and interpret data to generate insights.",
     "AI ინსტრუმენტებით მონაცემების შეგროვების, ანალიზის, ვიზუალიზაციისა და ინტერპრეტაციის უნარი.",
     "Data Analytics for Business (semester course)", "YES"],
    ["AI-ETHICS", "AI Ethics & Responsible Use", "AI ეთიკა და პასუხისმგებლიანი გამოყენება", "FOCUSED",
     "Understanding ethical considerations in AI use, including bias, transparency, privacy, accountability "
     "and responsible use.",
     "AI-ის გამოყენების ეთიკური ასპექტების ცოდნა: მიკერძოება, გამჭვირვალობა, კონფიდენციალურობა, ანგარიშვალდებულება.",
     "Responsible AI Masterclass", "YES"],
    ["AI-AGENTS", "AI Agents", "AI აგენტები", "APPLIED",
     "Ability to design and build AI agents that can perform tasks, make decisions and interact with tools "
     "and data.",
     "AI აგენტების დაპროექტებისა და აწყობის უნარი, რომლებიც ასრულებენ ამოცანებს და მუშაობენ ინსტრუმენტებთან.",
     "Build Your First AI Agent Challenge", "YES"],
    ["ML-FOUND", "Machine Learning Foundations", "მანქანური სწავლების საფუძვლები", "ADVANCED",
     "Ability to explain, train and evaluate core machine learning models on structured data.",
     "მანქანური სწავლების ძირითადი მოდელების ახსნის, გაწვრთნისა და შეფასების უნარი.",
     "Introduction to Machine Learning (semester course)", "YES"],
    ["AI-MARKETING", "AI for Marketing & Content", "AI მარკეტინგსა და კონტენტში", "FOCUSED",
     "Ability to use generative AI responsibly to plan, draft and evaluate marketing content.",
     "გენერაციული AI-ის პასუხისმგებლიანი გამოყენება მარკეტინგული კონტენტის დაგეგმვისა და შეფასებისთვის.",
     "AI in Marketing Guest Lecture", "YES"],
    ["CV-BASICS", "Computer Vision Basics", "კომპიუტერული ხედვის საფუძვლები", "APPLIED",
     "Ability to apply pre-trained vision models to classify and detect objects in images.",
     "წინასწარ გაწვრთნილი მოდელებით გამოსახულებებში ობიექტების კლასიფიკაციისა და აღმოჩენის უნარი.",
     "Computer Vision Hackathon", "YES"],
]


def _style_header(ws, tab):
    sys_cols = set(S.TABS[tab]["system"]) if tab in S.TABS else set()
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF" if cell.value not in sys_cols else "333333")
        cell.fill = SYS_FILL if cell.value in sys_cols else HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = max(14, min(48, len(str(cell.value)) + 6))
    ws.freeze_panes = "A2"


def _add_validation(ws, tab):
    hdr = S.headers(tab)

    def col(name):
        from openpyxl.utils import get_column_letter

        return get_column_letter(hdr.index(name) + 1)

    if tab == S.STUDENTS:
        dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
        dv.add(f"{col('consent')}2:{col('consent')}5000")
        ws.add_data_validation(dv)
    if tab == S.CATALOG:
        dv = DataValidation(type="list", formula1='"FOCUSED,APPLIED,ADVANCED"', allow_blank=False)
        dv.add(f"{col('stamp_type')}2:{col('stamp_type')}500")
        ws.add_data_validation(dv)
        dv2 = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
        dv2.add(f"{col('active')}2:{col('active')}500")
        ws.add_data_validation(dv2)
    if tab == S.AWARDS:
        dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
        dv.add(f"{col('confirmed')}2:{col('confirmed')}20000")
        ws.add_data_validation(dv)


def build_template(path: Path, *, students=None, awards=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = S.README
    for r in README_ROWS:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    for tab, rows in ((S.STUDENTS, students or []), (S.CATALOG, CATALOG_ROWS), (S.AWARDS, awards or [])):
        sheet = wb.create_sheet(tab)
        hdr = S.headers(tab)
        sheet.append(hdr)
        for r in rows:
            if isinstance(r, dict):
                sheet.append([r.get(h, "") for h in hdr])
            else:
                sheet.append(list(r) + [""] * (len(hdr) - len(r)))
        _style_header(sheet, tab)
        _add_validation(sheet, tab)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_csvs(xlsx: Path, out_dir: Path):
    wb = load_workbook(xlsx, read_only=True)
    for tab in (S.STUDENTS, S.CATALOG, S.AWARDS):
        with open(out_dir / f"{tab}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            for row in wb[tab].iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
    wb.close()


def read_csv_dicts(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))
