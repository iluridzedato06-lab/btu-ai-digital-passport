"""Local stand-in for Google Sheets: an .xlsx workbook with identical tabs and headers."""
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook

from .base import CellUpdate, SheetRow


class LocalXlsxAdapter:
    name = "local_xlsx"

    def __init__(self, path: Path | None = None):
        self.path = Path(path or settings.LOCAL_SHEET_PATH)

    def describe(self) -> str:
        return f"Local workbook {self.path.name}"

    def _ensure(self):
        if not self.path.exists():
            raise FileNotFoundError(
                f"Local sheet {self.path} not found. Run `python manage.py seed_demo` to create it."
            )

    def read_tab(self, tab: str) -> list[SheetRow]:
        self._ensure()
        wb = load_workbook(self.path, data_only=True, read_only=True)
        try:
            if tab not in wb.sheetnames:
                return []
            ws = wb[tab]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        if not rows:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for idx, raw in enumerate(rows[1:], start=2):
            values = {h: raw[i] if i < len(raw) else None for i, h in enumerate(header) if h}
            if all(v in (None, "") for v in values.values()):
                continue
            out.append(SheetRow(tab=tab, row_number=idx, values=values))
        return out

    def write_cells(self, updates: list[CellUpdate]) -> None:
        if not updates:
            return
        self._ensure()
        wb = load_workbook(self.path)
        try:
            for tab in {u.tab for u in updates}:
                ws = wb[tab]
                header = {str(c.value).strip(): c.column for c in ws[1] if c.value}
                for u in (x for x in updates if x.tab == tab):
                    col = header.get(u.column)
                    if col is None:
                        col = ws.max_column + 1
                        ws.cell(row=1, column=col, value=u.column)
                        header[u.column] = col
                    ws.cell(row=u.row_number, column=col, value=u.value)
            tmp = self.path.with_suffix(".tmp.xlsx")
            wb.save(tmp)
        finally:
            wb.close()
        tmp.replace(self.path)
